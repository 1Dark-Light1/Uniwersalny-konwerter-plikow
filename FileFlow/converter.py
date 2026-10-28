"""Logika konwersji oddzielona od interfejsu GUI."""
from __future__ import annotations

import csv
import json
import shutil
import subprocess
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

LOG = logging.getLogger("konwerter")
Converter = Callable[[Path, Path], Path]
REGISTRY: dict[tuple[str, str], Converter] = {}


def converter(source: str, target: str):
    """Rejestruje obsługiwany kierunek konwersji."""
    def register(func: Converter) -> Converter:
        REGISTRY[(source.lower(), target.lower())] = func
        return func
    return register


@dataclass(frozen=True)
class ConversionRequest:
    source: Path
    target_format: str
    output_dir: Path

    def validate(self) -> None:
        if not self.source.exists() or not self.source.is_file():
            raise FileNotFoundError(f"Nie znaleziono pliku wejściowego: {self.source}")
        if not self.output_dir.exists() or not self.output_dir.is_dir():
            raise NotADirectoryError(f"Katalog docelowy nie istnieje: {self.output_dir}")
        if self.source.suffix.lower() not in {src for src, _ in REGISTRY}:
            raise ValueError(f"Nieobsługiwane rozszerzenie pliku: {self.source.suffix or '(brak)'}")
        key = (self.source.suffix.lower(), self.target_format.lower())
        if key not in REGISTRY:
            raise ValueError(f"Nieobsługiwana konwersja: {key[0]} → {key[1]}")


def available_targets(source: Path) -> list[str]:
    """Zwraca wyłącznie formaty docelowe obsługiwane dla danego rozszerzenia."""
    return sorted(target for (src, target) in REGISTRY if src == source.suffix.lower())


def safe_output_path(output_dir: Path, stem: str, extension: str) -> Path:
    """Nie nadpisuje plików: dopisuje numer, jeśli nazwa już istnieje."""
    candidate = output_dir / f"{stem}{extension}"
    counter = 1
    while candidate.exists():
        candidate = output_dir / f"{stem}_{counter}{extension}"
        counter += 1
    return candidate


def _read_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Plik nie jest poprawnym tekstem UTF-8.") from exc


@converter(".csv", ".json")
def csv_to_json(source: Path, output_dir: Path) -> Path:
    with source.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            raise ValueError("CSV musi zawierać wiersz nagłówka.")
        rows = list(reader)
    target = safe_output_path(output_dir, source.stem, ".json")
    target.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    return target


@converter(".json", ".csv")
def json_to_csv(source: Path, output_dir: Path) -> Path:
    try:
        data = json.loads(_read_text(source))
    except json.JSONDecodeError as exc:
        raise ValueError(f"Niepoprawny JSON (linia {exc.lineno}, kolumna {exc.colno}).") from exc
    if not isinstance(data, list) or not data:
        raise ValueError("JSON musi być niepustą listą obiektów, np. [{\"imie\": \"Anna\"}].")
    if not all(isinstance(row, dict) for row in data):
        raise ValueError("Każdy element listy JSON musi być obiektem.")
    headers = list(dict.fromkeys(key for row in data for key in row.keys()))
    target = safe_output_path(output_dir, source.stem, ".csv")
    with target.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(data)
    return target


@converter(".txt", ".docx")
def txt_to_docx(source: Path, output_dir: Path) -> Path:
    text = _read_text(source)
    document = Document()
    for paragraph in text.splitlines():
        document.add_paragraph(paragraph)
    target = safe_output_path(output_dir, source.stem, ".docx")
    document.save(target)
    return target


@converter(".docx", ".txt")
def docx_to_txt(source: Path, output_dir: Path) -> Path:
    document = Document(source)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    target = safe_output_path(output_dir, source.stem, ".txt")
    target.write_text(text, encoding="utf-8")
    return target


@converter(".txt", ".pdf")
def txt_to_pdf(source: Path, output_dir: Path) -> Path:
    text = _read_text(source)
    target = safe_output_path(output_dir, source.stem, ".pdf")
    styles = getSampleStyleSheet()
    # Rejestrujemy popularny font Unicode, jeśli jest dostępny, aby zachować polskie znaki.
    font_candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/Library/Fonts/Arial.ttf"),
    ]
    font_name = "Helvetica"
    for font_path in font_candidates:
        if font_path.exists():
            try:
                pdfmetrics.registerFont(TTFont("UnicodeFont", str(font_path)))
                font_name = "UnicodeFont"
                break
            except Exception:
                pass
    styles["Normal"].fontName = font_name
    story = []
    for line in text.splitlines() or [""]:
        story.append(Paragraph(line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;") or " ", styles["Normal"]))
        story.append(Spacer(1, 6))
    SimpleDocTemplate(str(target), pagesize=A4).build(story)
    return target


def _libreoffice_path() -> str | None:
    for name in ("libreoffice", "soffice"):
        path = shutil.which(name)
        if path:
            return path
    candidates = [
        Path("C:/Program Files/LibreOffice/program/soffice.exe"),
        Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
    ]
    return next((str(p) for p in candidates if p.exists()), None)


@converter(".docx", ".pdf")
def docx_to_pdf(source: Path, output_dir: Path) -> Path:
    executable = _libreoffice_path()
    if not executable:
        raise RuntimeError("Nie znaleziono LibreOffice. Zainstaluj LibreOffice i uruchom aplikację ponownie.")
    # LibreOffice sam tworzy plik o nazwie źródła, dlatego używamy tymczasowego katalogu.
    import tempfile
    with tempfile.TemporaryDirectory(prefix="konwerter_") as temp:
        command = [executable, "--headless", "--convert-to", "pdf", "--outdir", temp, str(source)]
        try:
            result = subprocess.run(command, capture_output=True, text=True, timeout=90, check=False)
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError("LibreOffice przekroczył limit czasu 90 sekund.") from exc
        if result.returncode != 0:
            raise RuntimeError(result.stderr.strip() or result.stdout.strip() or "Konwersja DOCX → PDF nie powiodła się.")
        generated = Path(temp) / f"{source.stem}.pdf"
        if not generated.exists():
            raise RuntimeError("LibreOffice nie utworzył oczekiwanego pliku PDF.")
        target = safe_output_path(output_dir, source.stem, ".pdf")
        shutil.copy2(generated, target)
    return target


@converter(".pdf", ".docx")
def pdf_to_docx(source: Path, output_dir: Path) -> Path:
    """Eksportuje tekst PDF do DOCX; nie odtwarza oryginalnego układu ani grafiki."""
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise RuntimeError("Brak biblioteki PyMuPDF. Zainstaluj zależności z requirements.txt.") from exc
    doc = Document()
    try:
        with fitz.open(source) as pdf:
            for page in pdf:
                page_text = page.get_text("text")
                for paragraph in page_text.splitlines():
                    doc.add_paragraph(paragraph)
    except Exception as exc:
        raise ValueError(f"Nie udało się odczytać PDF: {exc}") from exc
    target = safe_output_path(output_dir, source.stem, ".docx")
    doc.save(target)
    return target


def execute(request: ConversionRequest) -> Path:
    request.validate()
    key = (request.source.suffix.lower(), request.target_format.lower())
    LOG.info("Konwersja rozpoczęta: %s → %s", request.source, request.target_format)
    try:
        result = REGISTRY[key](request.source, request.output_dir)
        LOG.info("Konwersja zakończona: %s", result)
        return result
    except Exception:
        LOG.exception("Błąd konwersji %s → %s", *key)
        raise
