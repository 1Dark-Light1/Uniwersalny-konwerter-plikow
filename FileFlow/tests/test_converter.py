import csv
import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from converter import (
    csv_to_json,
    json_to_csv,
    txt_to_docx,
    docx_to_txt,
    txt_to_pdf,
    docx_to_pdf,
    pdf_to_docx,
)


class TestConverter(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.temp_path = Path(self.temp_dir.name)
        self.output_dir = self.temp_path / "output"
        self.output_dir.mkdir()
        self.input_file = self.temp_path / "input.txt"

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_csv_to_json(self):
        csv_file = self.temp_path / "data.csv"

        with open(csv_file, "w", encoding="utf-8", newline="") as file:
            writer = csv.writer(file)
            writer.writerow(["name", "age"])
            writer.writerow(["Anna", "20"])

        json_file = csv_to_json(csv_file, self.output_dir)

        with open(json_file, "r", encoding="utf-8") as file:
            data = json.load(file)

        self.assertEqual(data[0]["name"], "Anna")
        self.assertEqual(data[0]["age"], "20")

    def test_json_to_csv(self):
        json_file = self.temp_path / "data.json"

        with open(json_file, "w", encoding="utf-8") as file:
            json.dump([{"name": "Anna", "age": 20}], file)

        csv_file = json_to_csv(json_file, self.output_dir)

        with open(csv_file, "r", encoding="utf-8-sig") as file:
            data = list(csv.DictReader(file))

        self.assertEqual(data[0]["name"], "Anna")
        self.assertEqual(data[0]["age"], "20")

    def test_txt_to_docx(self):
        from docx import Document

        with open(self.input_file, "w", encoding="utf-8") as file:
            file.write("Test document")

        output_docx = txt_to_docx(self.input_file, self.output_dir)

        document = Document(output_docx)
        self.assertIn("Test document", document.paragraphs[0].text)

    def test_docx_to_txt(self):
        from docx import Document

        docx_file = self.temp_path / "input.docx"
        document = Document()
        document.add_paragraph("Test document")
        document.save(docx_file)

        output_txt = docx_to_txt(docx_file, self.output_dir)

        with open(output_txt, "r", encoding="utf-8") as file:
            content = file.read()

        self.assertIn("Test document", content)

    def test_txt_to_pdf(self):
        with open(self.input_file, "w", encoding="utf-8") as file:
            file.write("Test PDF document")

        output_pdf = txt_to_pdf(self.input_file, self.output_dir)

        self.assertTrue(output_pdf.exists())
        self.assertGreater(output_pdf.stat().st_size, 0)

    def test_docx_to_pdf(self):
        from docx import Document

        docx_file = self.temp_path / "input.docx"
        document = Document()
        document.add_paragraph("Test PDF document")
        document.save(docx_file)

        output_pdf = docx_to_pdf(docx_file, self.output_dir)

        self.assertTrue(output_pdf.exists())
        self.assertGreater(output_pdf.stat().st_size, 0)
        import fitz

        with fitz.open(output_pdf) as pdf:
            self.assertIn("Test PDF document", pdf[0].get_text())

    def test_pdf_to_docx(self):
        with open(self.input_file, "w", encoding="utf-8") as file:
            file.write("Test PDF document")
        pdf_file = txt_to_pdf(self.input_file, self.temp_path)

        output_docx = pdf_to_docx(pdf_file, self.output_dir)

        self.assertTrue(output_docx.exists())
        self.assertGreater(output_docx.stat().st_size, 0)
        from docx import Document

        self.assertIn("Test PDF document", Document(output_docx).paragraphs[0].text)


if __name__ == "__main__":
    unittest.main()